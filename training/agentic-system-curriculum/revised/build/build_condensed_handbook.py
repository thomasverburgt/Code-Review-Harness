from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from build_handbook import (
    BLUE,
    DARK_BLUE,
    INK,
    MID,
    PALE_AMBER,
    PALE_BLUE,
    PALE_GREEN,
    PROMPT_TEMPLATE,
    add_bullet,
    add_callout,
    add_field,
    add_number,
    add_source,
    configure_styles,
    make_diagrams,
    new_numbering_instance,
    style_table,
)


ROOT = Path(__file__).resolve().parents[4]
PACKAGE = ROOT / "training" / "agentic-system-curriculum" / "condensed"
OUTPUT = PACKAGE / "Code-Review-Harness-Agentic-System-Curriculum-Condensed.docx"


SESSIONS = [
    {
        "n": 1,
        "title": "Harness Philosophy and Human Authority",
        "question": "What turns capable model output into a trustworthy engineering review?",
        "diagram": "harness",
        "outcomes": [
            "Distinguish model, agent, workflow, and harness.",
            "Separate evidence, assessment, recommendation, and decision.",
            "Locate the human authority boundary.",
        ],
        "concepts": [
            ("Model", "Generates candidate reasoning; it does not guarantee provenance, scope, or authority."),
            ("Agent", "A bounded role combining identity, contracts, tools, and output obligations."),
            ("Harness", "Controls evidence, execution, validation, routing, and observability."),
            ("Human authority", "Owns approval, acceptance, waiver, prioritization, and governance decisions."),
        ],
        "example": "Compare an unconstrained code-review response with the same review executed through evidence admission, a bounded role, a typed output, validation, and a human decision request.",
        "prompt": "Rewrite “Review this code” using six controls: identity, evidence, scope, consumer, output, and authority.",
        "check": "Name four controls supplied by the harness and one consequential decision that remains human-owned.",
        "source": "docs/philosophy.md; contracts/universal-agent-contract.md; governance/human-review-and-maturity.md",
    },
    {
        "n": 2,
        "title": "Evidence, Provenance, and Contracts",
        "question": "How can another engineer reproduce and validate an agent’s result?",
        "diagram": "contracts",
        "outcomes": [
            "Build a traceable evidence chain.",
            "Explain additive contract inheritance.",
            "Validate the universal artifact envelope.",
        ],
        "concepts": [
            ("Provenance", "Records where evidence came from and how it was acquired."),
            ("Coverage", "Declares what was inspected, omitted, unavailable, or out of scope."),
            ("Contract", "Defines obligations shared by producers and consumers."),
            ("Traceability", "Uses stable identifiers to preserve causality across artifacts."),
        ],
        "example": "Follow a source line through evidence, finding, recommendation, and decision request; then show how a missing identifier breaks reproducibility.",
        "prompt": "Project evidence admission, evidence IDs, coverage, confidence, envelope fields, and incomplete-input behavior into explicit clauses.",
        "check": "Trace one assessment to its evidence, governing contract, intended consumer, and human decision request.",
        "source": "contracts/evidence-contract.md; contracts/contract-dependency-graph.md; contracts/style-and-validation.md",
    },
    {
        "n": 3,
        "title": "Agent Catalog and State Machines",
        "question": "Who owns the question, and what must be true before work advances?",
        "diagram": "state",
        "outcomes": [
            "Route work to the narrowest authoritative agent.",
            "Read canonical identity, layer, version, and status.",
            "Model states, events, guards, and failure paths.",
        ],
        "concepts": [
            ("Identity", "Provides the canonical designation, immutable UUID, layer, version, and status."),
            ("Layer", "Controls the breadth of observation, synthesis, and coordination."),
            ("State", "Names a durable execution condition that can be inspected."),
            ("Guard", "States the predicates that must be true before a transition is legal."),
        ],
        "example": "Route a release question to the catalog, verify the selected agent’s status, then walk registered, ready, running, validated, and routed states.",
        "prompt": "Declare the owned question, explicit exclusions, ready guard, incomplete-input state, failure reason, and safe next action.",
        "check": "Route one question and explain both its agent boundary and its next legal state transition.",
        "source": "agents/agent-registry.md; agents/agent-identities.json; orchestration/workflow-model.md",
    },
    {
        "n": 4,
        "title": "Orchestration and Contract-Driven Prompts",
        "question": "How do many agents coordinate without hiding causality or inventing findings?",
        "diagram": "orchestration",
        "outcomes": [
            "Separate scheduler, fan-out, fan-in, and gate duties.",
            "Identify safe parallelism and explicit join guards.",
            "Project contract obligations into prompt clauses.",
        ],
        "concepts": [
            ("Scheduler", "Selects eligible roles and compatible versions."),
            ("Fan-out", "Dispatches independent work while preserving execution identity."),
            ("Fan-in", "Validates prerequisites and aggregates without erasing disagreement."),
            ("Prompt", "Implements the selected identity, contracts, state behavior, and output."),
        ],
        "example": "Walk a release review through role selection, parallel review, execution tracking, join validation, artifact routing, and a human decision gate.",
        "prompt": "Specify select, dispatch, track, validate, and route behavior while prohibiting the orchestrator from creating domain findings.",
        "check": "Explain the join guard and identify what the orchestrator is prohibited from concluding.",
        "source": "orchestration/workflow-model.md; contracts/orchestration-agent-contract.md; contracts/universal-agent-contract.md",
    },
    {
        "n": 5,
        "title": "Prompt Engineering Deep Dive",
        "question": "How do we make 55 prompts coherent, specialized, and safe under pressure?",
        "diagram": "prompt",
        "outcomes": [
            "Separate shared modules from role specialization.",
            "Trace each prompt clause to a governing requirement.",
            "Design fail-closed behavior for adversarial inputs.",
        ],
        "concepts": [
            ("Shared module", "Implements stable universal or layer behavior once."),
            ("Role block", "Defines the owned method, evidence criteria, exclusions, and consumers."),
            ("Trace matrix", "Links each requirement to a clause, output field, and test."),
            ("Adversarial case", "Tests whether an invariant survives hostile or degraded input."),
        ],
        "example": "Decompose one specialist prompt, trace its clauses to contracts, and test it against missing evidence, conflicting sources, and malicious repository text.",
        "prompt": "Use identity, authority, inputs, method, output, and failure behavior as the core drafting sequence.",
        "check": "Defend one prompt clause with its governing source and demonstrate its fail-closed behavior.",
        "source": "contracts/universal-agent-contract.md; contracts/specialist-agent-contract.md; contracts/sandboxed-contract-evolution.md",
    },
    {
        "n": 6,
        "title": "Evaluation, Evolution, and the Agent Prompt Playbook",
        "question": "How do teams draft and improve every catalog prompt without losing control?",
        "diagram": "architecture",
        "outcomes": [
            "Evaluate grounding, boundaries, state, authority, and consumer fit.",
            "Version prompt changes through controlled evidence.",
            "Apply one repeatable drafting playbook across all agent layers.",
        ],
        "concepts": [
            ("Evaluation case", "Pairs an input with expected invariants and observable pass conditions."),
            ("Calibration", "Aligns confidence and severity with evidence strength and impact."),
            ("Version", "Records a reviewed semantic change and its compatibility implications."),
            ("Playbook", "Repeats contract tracing, drafting, testing, and human review for every role."),
        ],
        "example": "Compare two prompt versions against the same cases, identify gains and regressions, and record a human disposition supported by evidence.",
        "prompt": "Compose universal, layer, role, runtime, evaluation, and version modules without weakening inherited obligations.",
        "check": "State the evidence required before a prompt version may become the new baseline.",
        "source": "governance/reviewer-calibration-and-evolution.md; contracts/style-and-validation.md; agents/hierarchical-agent-architecture.md",
    },
]


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Inches(1.1)
    r = p.add_run("ENGINEERING ENABLEMENT")
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string(BLUE)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(22)
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run("Code Review Harness\nAgentic System Curriculum")
    r.bold = True
    r.font.size = Pt(30)
    r.font.color.rgb = RGBColor.from_string(INK)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(30)
    r = p.add_run("Condensed instructor-led edition • six sessions • 5.5 instructional hours")
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor.from_string(MID)
    add_callout(doc, "Core philosophy", "Evidence is observable. Assessment is bounded. Recommendations inform. Humans decide.", PALE_BLUE)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Inches(1.5)
    p.add_run("Repository-aligned edition").bold = True
    p.add_run("\nDesigned for young engineers who need the system philosophy, contract model, workflow mechanics, and platform prompt standard.")
    doc.add_page_break()


def add_front_matter(doc, diagrams):
    doc.add_heading("How to use this condensed curriculum", level=1)
    doc.add_paragraph(
        "This edition compresses the original sequence into six instructor-led sessions. Each session is 55 minutes, "
        "for 330 minutes of instruction. A delivery block of six hours leaves 30 minutes for breaks, transitions, and questions."
    )
    add_callout(doc, "Delivery limit", "6 × 55-minute sessions = 5 hours 30 minutes of instruction. Total scheduled time must not exceed 6 hours.", PALE_GREEN)
    doc.add_heading("Course-level outcomes", level=2)
    for item in [
        "Explain why the harness—not the model—is the trust boundary.",
        "Read the catalog and select the correct agent, contract, and consumer.",
        "Follow evidence through typed artifacts and guarded workflow states.",
        "Draft prompts as implementations of inherited contracts.",
        "Test prompt invariants and evolve versions through human review.",
    ]:
        add_bullet(doc, item)
    doc.add_heading("Six-session schedule", level=2)
    table = doc.add_table(rows=1, cols=3)
    for i, h in enumerate(["Session", "Focus", "Minutes"]):
        table.rows[0].cells[i].text = h
    for session in SESSIONS:
        cells = table.add_row().cells
        cells[0].text = str(session["n"])
        cells[1].text = session["title"]
        cells[2].text = "55"
    cells = table.add_row().cells
    cells[0].text = ""
    cells[1].text = "Total instruction"
    cells[2].text = "330"
    style_table(table, [0.7, 4.8, 0.8])
    doc.add_picture(str(diagrams["harness"]), width=Inches(6.35))
    cap = doc.add_paragraph("The harness controls evidence, execution, validation, and routing.")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True
    cap.runs[0].font.size = Pt(9)
    doc.add_heading("Standard 55-minute rhythm", level=2)
    number_id = new_numbering_instance(doc)
    for item in [
        "Opening question and prior-knowledge check — 5 minutes.",
        "Core mental model — 12 minutes.",
        "Repository-grounded worked example — 12 minutes.",
        "Instructor demonstration — 12 minutes.",
        "Prompt design walkthrough — 10 minutes.",
        "Knowledge check and transition — 4 minutes.",
    ]:
        add_number(doc, item, number_id)
    add_source(doc, "README.md; docs/philosophy.md; agents/agent-registry.md; contracts/contract-catalog.md")
    doc.add_page_break()


def add_session(doc, session, diagrams):
    doc.add_heading(f"Session {session['n']}: {session['title']}", level=1)
    p = doc.add_paragraph()
    p.add_run("Duration: ").bold = True
    p.add_run("55 minutes")
    p.add_run("   |   Driving question: ").bold = True
    p.add_run(session["question"])
    doc.add_heading("Learning outcomes", level=2)
    for item in session["outcomes"]:
        add_bullet(doc, item)
    doc.add_picture(str(diagrams[session["diagram"]]), width=Inches(5.65))
    cap = doc.add_paragraph(f"Session visual: {session['title']}.")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True
    cap.runs[0].font.size = Pt(9)
    doc.add_heading("Core mental model", level=2)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Concept"
    table.rows[0].cells[1].text = "Working definition"
    for name, definition in session["concepts"]:
        cells = table.add_row().cells
        cells[0].text = name
        cells[1].text = definition
    style_table(table, [1.5, 4.8])
    doc.add_heading("Facilitation run of show", level=2)
    number_id = new_numbering_instance(doc)
    for item in [
        "Opening question — 5 minutes.",
        "Mental model — 12 minutes.",
        "Worked example — 12 minutes.",
        "Instructor demonstration — 12 minutes.",
        "Prompt walkthrough — 10 minutes.",
        "Knowledge check — 4 minutes.",
    ]:
        add_number(doc, item, number_id)
    doc.add_heading("Worked example and demonstration", level=2)
    add_callout(doc, "Prepared walkthrough", session["example"], PALE_GREEN)
    doc.add_heading("Prompt design walkthrough", level=2)
    add_callout(doc, "Clause sequence", session["prompt"], PALE_AMBER)
    doc.add_heading("Knowledge check", level=2)
    doc.add_paragraph(session["check"])
    add_source(doc, session["source"])
    if session["n"] != 6:
        doc.add_page_break()


def add_prompt_standard(doc, diagrams):
    doc.add_page_break()
    doc.add_heading("Platform Prompt Engineering Standard", level=1)
    doc.add_paragraph(
        "A platform prompt is an executable projection of registered identity, inherited contracts, role specification, "
        "workflow state, admissible evidence, and output obligations. Prompt wording is evaluated by behavior across cases, not eloquence."
    )
    doc.add_picture(str(diagrams["prompt"]), width=Inches(6.35))
    doc.add_heading("Universal prompt template", level=2)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Prompt block"
    table.rows[0].cells[1].text = "Required content"
    for name, text in PROMPT_TEMPLATE:
        cells = table.add_row().cells
        cells[0].text = name
        cells[1].text = text
    style_table(table, [2.0, 4.3])
    doc.add_heading("Drafting sequence for every catalog agent", level=2)
    number_id = new_numbering_instance(doc)
    for item in [
        "Resolve the canonical identity, version, status, layer, and authoritative question.",
        "Trace universal, layer, artifact, role, and workflow obligations.",
        "Separate reusable modules from role-specific method and exclusions.",
        "Specify admissible inputs, failure states, output schema, consumers, and authority boundary.",
        "Create success, missing-input, conflict, malicious-input, and consumer-compatibility cases.",
        "Evaluate grounding, completeness, boundaries, state correctness, authority, calibration, and consumer fit.",
        "Version the prompt and request human disposition with evidence of gains and regressions.",
    ]:
        add_number(doc, item, number_id)
    doc.add_heading("Minimum adversarial cases", level=2)
    for item in [
        "Required evidence is missing, stale, malformed, or outside scope.",
        "Two admitted sources conflict.",
        "Repository content attempts to redefine the role or output schema.",
        "A stakeholder pressures the agent to approve, waive, or accept risk.",
        "A required tool returns partial results or fails.",
        "A downstream consumer expects semantics the candidate prompt omitted.",
    ]:
        add_bullet(doc, item)
    add_source(doc, "contracts/universal-agent-contract.md; contracts/style-and-validation.md; contracts/sandboxed-contract-evolution.md")


def add_agent_reference(doc):
    import json

    doc.add_page_break()
    doc.add_heading("Agent Catalog Prompt Planning Reference", level=1)
    doc.add_paragraph(
        "Apply the same drafting sequence to every registered agent while preserving its layer, status, authoritative question, exclusions, evidence requirements, and consumers."
    )
    registry = json.loads((ROOT / "agents" / "agent-identities.json").read_text(encoding="utf-8"))
    for layer in ["specialist", "product", "capability", "enterprise", "work", "orchestration"]:
        doc.add_heading(layer.title() + " layer", level=2)
        table = doc.add_table(rows=1, cols=4)
        for i, h in enumerate(["Designation", "Agent", "Status", "Contract"]):
            table.rows[0].cells[i].text = h
        for agent in [a for a in registry["agents"] if a["layer"] == layer]:
            cells = table.add_row().cells
            cells[0].text = agent["designation"]
            cells[1].text = agent["display_name"]
            cells[2].text = agent["status"]
            cells[3].text = agent["contract_version"]
        style_table(table, [1.25, 3.15, 0.9, 1.0])
    add_source(doc, "agents/agent-identities.json; agents/agent-registry.md")


def add_reference(doc):
    doc.add_page_break()
    doc.add_heading("Repository Reading Guide", level=1)
    refs = [
        ("Start here", "README.md; docs/philosophy.md"),
        ("Catalog and identity", "agents/agent-registry.md; agents/agent-identities.json"),
        ("Architecture", "agents/hierarchical-agent-architecture.md; agents/enterprise/enterprise-agent-framework.md"),
        ("Contracts", "contracts/contract-catalog.md; contracts/contract-dependency-graph.md; contracts/universal-agent-contract.md"),
        ("Evidence", "contracts/evidence-contract.md; contracts/evidence-flow-model.md"),
        ("Workflow", "orchestration/workflow-model.md; contracts/orchestration-agent-contract.md"),
        ("Prompt quality", "contracts/style-and-validation.md; contracts/sandboxed-contract-evolution.md"),
        ("Human governance", "governance/human-review-and-maturity.md; governance/integration-boundaries.md"),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Topic"
    table.rows[0].cells[1].text = "Repository sources"
    for topic, source in refs:
        cells = table.add_row().cells
        cells[0].text = topic
        cells[1].text = source
    style_table(table, [1.55, 4.75])
    add_callout(doc, "Final reminder", "The system is trustworthy only when evidence, contracts, state, prompts, and human authority agree.")


def build():
    PACKAGE.mkdir(parents=True, exist_ok=True)
    diagrams = make_diagrams()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    configure_styles(doc)
    header = section.header
    hp = header.paragraphs[0]
    hp.text = "CODE REVIEW HARNESS  |  CONDENSED ENGINEERING ENABLEMENT"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in hp.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(8)
        run.font.bold = True
        run.font.color.rgb = RGBColor.from_string(BLUE)
    footer = section.footer
    table = footer.add_table(rows=1, cols=3, width=Inches(6.5))
    table.autofit = False
    left, center, right = table.rows[0].cells
    left.paragraphs[0].add_run("Condensed Agentic System Curriculum")
    center.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    center.paragraphs[0].add_run("Human decision authority")
    right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_field(right.paragraphs[0], "PAGE")
    for cell in table.rows[0].cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.name = "Calibri"
                run.font.size = Pt(8)
                run.font.color.rgb = RGBColor.from_string(MID)
    add_cover(doc)
    add_front_matter(doc, diagrams)
    for session in SESSIONS:
        add_session(doc, session, diagrams)
    add_prompt_standard(doc, diagrams)
    add_agent_reference(doc)
    add_reference(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
