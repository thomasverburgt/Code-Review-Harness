#!/usr/bin/env python3
"""Verify specialist Word handouts reproduce the bound 0.2.0 prompts exactly."""

from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "reports/specialist-human-review-handouts/2026-08-04"
PROMPT_MANIFEST = ROOT / "appendices/prompt-templates/candidates/manifest-design-0.2.0.json"
BANNED = (
    r"\breviewer track\b",
    r"\bdomain qualifications\b",
    r"\bauthorized classifications\b",
    r"\bconflict of interest\b",
    r"\bstarted at\b",
    r"\binterns?\b",
    r"design-0\.1\.0 prompt",
)


def load(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def main() -> int:
    prompt_manifest = load(PROMPT_MANIFEST)
    handout_manifest = load(OUTPUT / "manifest.json")
    assert len(prompt_manifest["prompts"]) == 22
    assert len(handout_manifest["handouts"]) == 22

    checked: set[str] = set()
    for entry in handout_manifest["handouts"]:
        designation = entry["designation"]
        assert designation not in checked
        checked.add(designation)
        assert entry["prompt"]["version"] == "design-0.2.0"

        docx_path = ROOT / entry["document"]
        assert docx_path.is_file()
        assert entry["sha256"] == "sha256:" + hashlib.sha256(docx_path.read_bytes()).hexdigest()
        document = Document(docx_path)
        full_text = "\n".join(
            [paragraph.text for paragraph in document.paragraphs]
            + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        )
        for pattern in BANNED:
            assert not re.search(pattern, full_text, flags=re.IGNORECASE), f"{designation}: banned text {pattern}"
        assert "design-0.2.0" in full_text
        assert "Appendix A - Prompt record and markup copy" in full_text

        prompt_path = ROOT / prompt_manifest["prompts"][designation]["path"]
        prompt_lines = prompt_path.read_text(encoding="utf-8").splitlines()
        prompt_table = document.tables[-1]
        assert [cell.text for cell in prompt_table.rows[0].cells] == ["Line", "Prompt text"]
        reproduced = [row.cells[1].text for row in prompt_table.rows[1:]]
        numbering = [row.cells[0].text for row in prompt_table.rows[1:]]
        assert reproduced == prompt_lines, f"{designation}: prompt appendix differs from source"
        assert numbering == [str(index) for index in range(1, len(prompt_lines) + 1)]

    assert checked == set(prompt_manifest["prompts"])
    print(json.dumps({"handouts": len(checked), "prompt_version": "design-0.2.0", "exact_prompt_copies": True}))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
