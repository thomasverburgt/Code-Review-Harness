#!/usr/bin/env python3
"""Build one visually consistent Word human-review handout per ADR-0039 specialist."""
from __future__ import annotations

import hashlib
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "reports" / "specialist-human-review-handouts" / "2026-08-04"
DATE = "August 4, 2026"
BLUE = "2E74B5"
NAVY = "1F4D78"
LIGHT_BLUE = "E8EEF5"
PALE_BLUE = "F3F7FA"
GRAY = "666666"
LIGHT_GRAY = "F3F4F6"
RED = "9C0006"
PALE_RED = "FCE8E6"
GREEN = "276749"
WHITE = "FFFFFF"
BLACK = "000000"
USABLE_DXA = 9360

PACKET_ROOTS = [
    ROOT / "fixtures/specialist-static-evidence-increment1/evidence/2026-08-03/adr0039/gx10-live-run-001/live-run",
    ROOT / "fixtures/specialist-build-supply-chain-increment2/evidence/2026-08-03/adr0039/gx10-live-run-001/live-run",
    ROOT / "fixtures/specialist-k8s-secure-increment34/evidence/2026-08-03/adr0039/gx10-live-run-001/live-run",
    ROOT / "fixtures/specialist-increment567/evidence/2026-08-03/adr0039/gx10-live-run-001",
    ROOT / "fixtures/specialist-calibration-0.2.0/evidence/2026-08-04/adr0039/gx10-live-run-001",
]

QUALIFIED_SEMANTIC = {"SPEC-ARCH", "SPEC-DATA", "SPEC-DIAGRAM", "SPEC-RESEARCH", "SPEC-RISK", "SPEC-SECURITY"}


def load(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def set_run(run, size=11, bold=False, italic=False, color=BLACK, font="Calibri"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    return run


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths[index]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_margins(cell)


def no_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "nil")


def bottom_border(paragraph, color="A6A6A6"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, NAVY, 10, 5),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = instruction
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t"); text.text = "1"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def configure_page(doc, designation, classification):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.32)
    section.footer_distance = Inches(0.32)
    header = section.header
    table = header.add_table(rows=1, cols=2, width=Inches(6.5))
    set_table_geometry(table, [4680, 4680]); no_table_borders(table)
    left = table.cell(0, 0).paragraphs[0]
    left.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run(left.add_run("ADR-0039 HUMAN SHADOW REVIEW"), 8.5, True, color=GRAY)
    right = table.cell(0, 1).paragraphs[0]
    right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    color = RED if classification != "public" else GRAY
    set_run(right.add_run(f"{designation} | {classification.upper()}"), 8.5, True, color=color)
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("Comparison-only calibration material  |  Page "), 8.5, color=GRAY)
    add_field(p, "PAGE")


def add_title_block(doc, profile, packet):
    designation = packet["candidate"]["designation"]
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(2)
    set_run(p.add_run("SPECIALIST REVIEW HANDOUT"), 10, True, color=BLUE)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    set_run(p.add_run(profile["agent"]["display_name"]), 24, True, color=NAVY)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    set_run(p.add_run(f"{designation} | Non-authoritative human shadow review"), 13, color=GRAY)
    restricted = packet["classification"] != "public"
    callout = doc.add_table(rows=1, cols=1)
    set_table_geometry(callout, [USABLE_DXA])
    cell = callout.cell(0, 0)
    shade(cell, PALE_RED if restricted else PALE_BLUE)
    p = cell.paragraphs[0]
    label = "RESTRICTED HANDLING" if restricted else "PURPOSE AND AUTHORITY"
    set_run(p.add_run(label + "\n"), 10, True, color=RED if restricted else NAVY)
    text = ("Do not distribute to general engineers. Restricted or qualified reviewers only. The matched value is redacted and must not be requested, copied, or reproduced in this handout. " if restricted else "This handout may be used by an engineer for traceability, clarity, reproducibility, and apparent support review. ")
    text += "Completion creates calibration evidence only. It does not approve the agent, establish professional truth, accept risk, authorize a change, schedule work, permit report use, or authorize deployment."
    set_run(p.add_run(text), 10.5, color=BLACK)
    doc.add_paragraph()
    metadata = [
        ("Packet", packet["packet_id"]),
        ("Packet hash", packet["packet_hash"]),
        ("Candidate artifact", packet["candidate_artifact_id"]),
        ("Focus binding", f"{packet['focus_binding']['profile_id']} v{packet['focus_binding']['profile_version']}"),
        ("Classification", packet["classification"]),
        ("Prepared", DATE),
    ]
    table = doc.add_table(rows=len(metadata), cols=2)
    set_table_geometry(table, [2100, 7260])
    for row, (label, value) in zip(table.rows, metadata):
        shade(row.cells[0], LIGHT_BLUE)
        set_run(row.cells[0].paragraphs[0].add_run(label), 9.5, True, color=NAVY)
        set_run(row.cells[1].paragraphs[0].add_run(str(value)), 9.5)
    doc.add_heading("Reviewer information", level=2)
    fields = ["Reviewer ID / name", "Experience level", "Completed at"]
    for label in fields:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        set_run(p.add_run(label + ": "), 9.5, True, color=NAVY)
        set_run(p.add_run(" "), 9.5)
        bottom_border(p)


def add_instructions_and_focus(doc, profile, packet):
    doc.add_page_break()
    doc.add_heading("How to complete this review", level=1)
    steps = [
        "Confirm that you are permitted to access the packet classification.",
        "For every item, follow the evidence locator to the exact immutable revision, path, line or section, and safe excerpt.",
        "Judge whether the statement is supported within the admitted evidence population. Do not fill evidence gaps with assumptions.",
        "Record one disposition, explain your rationale, assess locator correctness, and identify missing evidence.",
        "Complete the overall usability section and acknowledge the non-authoritative boundary before returning the handout.",
    ]
    for step in steps:
        p = doc.add_paragraph(style="List Number")
        p.add_run(step)
    doc.add_heading("Disposition guide", level=2)
    dispositions = [
        ("Supported", "The statement is fully supported by the admitted evidence."),
        ("Partially supported", "Some material is supported, but part is missing, overstated, or ambiguous."),
        ("Unsupported", "The admitted evidence does not support the statement."),
        ("Duplicate", "The statement materially duplicates another packet item."),
        ("Outside specialist scope", "The statement exceeds this agent's authorized focus."),
        ("Unable to determine", "The evidence is inaccessible, insufficient, unclear, or requires expertise you do not have."),
    ]
    table = doc.add_table(rows=1, cols=2)
    set_table_geometry(table, [2400, 6960])
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    for cell, text in zip(table.rows[0].cells, ("Disposition", "Use when")):
        shade(cell, LIGHT_BLUE); set_run(cell.paragraphs[0].add_run(text), 9.5, True, color=NAVY)
    for label, meaning in dispositions:
        cells = table.add_row().cells
        set_run(cells[0].paragraphs[0].add_run(label), 9.5, True)
        set_run(cells[1].paragraphs[0].add_run(meaning), 9.5)
        for cell in cells: cell_margins(cell)
    doc.add_heading("Agent focus", level=1)
    p = doc.add_paragraph()
    set_run(p.add_run("Mission. "), 10.5, True, color=NAVY)
    set_run(p.add_run(profile["focus"]["mission"]), 10.5)
    doc.add_heading("Authoritative question", level=2)
    for question in profile["focus"]["authoritative_questions"]:
        p = doc.add_paragraph(style="List Bullet"); p.add_run(question)
    doc.add_heading("Domain dimensions", level=2)
    for dimension in profile["focus"]["dimensions"]:
        p = doc.add_paragraph(style="List Bullet"); p.add_run(dimension)
    if packet["candidate"]["designation"] in QUALIFIED_SEMANTIC:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        set_run(p.add_run("ENGINEER LIMIT: "), 10, True, color=RED)
        set_run(p.add_run("An engineer may evaluate traceability, clarity, reproducibility, scope discipline, and apparent support. A qualified subject-matter expert must adjudicate the domain meaning or conclusion."), 10)


def unique_locators(packet):
    result = {}
    for item in packet["items"]:
        for locator in item["locators"]:
            result[locator["locator_id"]] = locator
    return list(result.values())


def add_evidence_map(doc, packet):
    doc.add_page_break()
    doc.add_heading("Evidence map", level=1)
    p = doc.add_paragraph()
    set_run(p.add_run("Evidence rule. "), 10.5, True, color=NAVY)
    set_run(p.add_run("Use only the evidence below. A correct path is not enough: verify the immutable revision, exact line or section, safe excerpt, and fingerprint. Record any mismatch in the item review."), 10.5)
    for index, loc in enumerate(unique_locators(packet), 1):
        doc.add_heading(f"Evidence {index}: {loc['evidence_id']}", level=2)
        table = doc.add_table(rows=0, cols=2)
        set_table_geometry(table, [2100, 7260])
        entries = [
            ("Repository", loc["repository_uri"]),
            ("Revision", loc["immutable_revision"]),
            ("Location", f"{loc['path']} | lines {loc['line_start']}-{loc['line_end']} | {loc['symbol_or_section']}"),
            ("Classification", loc["classification"]),
            ("Redaction", "applied; raw value omitted" if loc["redaction_applied"] else "none; raw restricted value not included"),
            ("Fingerprint", loc["line_fingerprint"]),
        ]
        for label, value in entries:
            cells = table.add_row().cells
            shade(cells[0], LIGHT_BLUE)
            set_run(cells[0].paragraphs[0].add_run(label), 9, True, color=NAVY)
            set_run(cells[1].paragraphs[0].add_run(str(value)), 9)
            for cell in cells: cell_margins(cell)
        p = doc.add_paragraph()
        set_run(p.add_run("Safe excerpt: "), 9.5, True, color=NAVY)
        set_run(p.add_run(loc["safe_excerpt"]), 9.5, italic=True)
        p = doc.add_paragraph()
        set_run(p.add_run("Reproduce: "), 9.5, True, color=NAVY)
        set_run(p.add_run(" ".join(loc["reproduction_steps"])), 9.5)


def add_response_lines(doc, label, count=3):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    set_run(p.add_run(label), 9.5, True, color=NAVY)
    for _ in range(count):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        set_run(p.add_run(" "), 10)
        bottom_border(p)


def option_line(doc, label, options):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(5)
    set_run(p.add_run(label + " "), 9.5, True, color=NAVY)
    set_run(p.add_run("   ".join(f"[ ] {option}" for option in options)), 9.5)


def add_item_review(doc, item, index):
    doc.add_page_break()
    doc.add_heading(f"Item {index} review", level=1)
    meta = doc.add_table(rows=3, cols=2)
    set_table_geometry(meta, [2100, 7260])
    for row, (label, value) in zip(meta.rows, (("Item", f"{item['item_id']} | {item['kind']}"), ("Source record", item["source_record_id"]), ("Evidence refs", ", ".join(item["evidence_refs"])) )):
        shade(row.cells[0], LIGHT_BLUE)
        set_run(row.cells[0].paragraphs[0].add_run(label), 9.5, True, color=NAVY)
        set_run(row.cells[1].paragraphs[0].add_run(value), 9.5)
    doc.add_heading("Statement under review", level=2)
    callout = doc.add_table(rows=1, cols=1)
    set_table_geometry(callout, [USABLE_DXA])
    shade(callout.cell(0, 0), PALE_BLUE)
    set_run(callout.cell(0, 0).paragraphs[0].add_run(item["statement"]), 11, True, color=BLACK)
    p = doc.add_paragraph()
    set_run(p.add_run(f"Agent confidence: {item['confidence']:.2f}. "), 9.5, True, color=NAVY)
    set_run(p.add_run(item["confidence_meaning"]), 9.5, italic=True, color=GRAY)
    doc.add_heading("Your assessment", level=2)
    option_line(doc, "Disposition:", ["supported", "partially supported", "unsupported", "duplicate", "outside scope", "unable to determine"])
    option_line(doc, "Locator correctness:", ["correct", "partially correct", "incorrect", "not verified"])
    option_line(doc, "Severity:", ["reasonable", "too high", "too low", "not applicable", "unable to determine"])
    option_line(doc, "Recommendation:", ["reasonable", "partially reasonable", "unreasonable", "not applicable", "unable to determine"])
    add_response_lines(doc, "Rationale - explain how the admitted evidence supports your selections:", 4)
    add_response_lines(doc, "Missing evidence - list what would change or strengthen your conclusion:", 3)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    set_run(p.add_run("Reviewer check: "), 9.5, True, color=NAVY)
    set_run(p.add_run("[ ] I inspected the exact evidence locator.   [ ] I did not use unstated assumptions.   [ ] I preserved unknowns."), 9.5)


def add_overall_review(doc, profile, packet):
    doc.add_page_break()
    doc.add_heading("Overall review and return record", level=1)
    option_line(doc, "Overall usability rating:", ["1 unusable", "2 weak", "3 usable with revision", "4 useful", "5 highly useful"])
    option_line(doc, "Would you request more evidence?", ["yes", "no"])
    add_response_lines(doc, "Overall usability rationale:", 5)
    add_response_lines(doc, "Consolidated evidence requests:", 5)
    doc.add_heading("Focus and boundary check", level=2)
    checks = [
        "The output materially addressed the named specialist focus rather than generic code review.",
        "Every actionable claim retained an exact source record and evidence locator.",
        "Missing evidence and uncertainty remained visible.",
        "The candidate stayed within its authority and did not imply approval, release, scheduling, deployment, or risk acceptance.",
        "Any disagreement or apparent duplicate was recorded rather than silently merged.",
    ]
    for check in checks:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        set_run(p.add_run("[ ] Yes   [ ] No   [ ] Unable to determine   "), 9.5, True, color=NAVY)
        set_run(p.add_run(check), 9.5)
    doc.add_heading("Non-authoritative reviewer acknowledgment", level=2)
    p = doc.add_paragraph()
    set_run(p.add_run("[ ] I acknowledge that this response is non-authoritative human review with calibration-evidence-only effect. It does not establish professional truth, approve the agent, accept risk, determine compliance, authorize a change, schedule work, permit report use, or authorize deployment."), 10, True if packet["classification"] != "public" else False)
    if packet["classification"] != "public":
        p = doc.add_paragraph()
        set_run(p.add_run("[ ] I am authorized for this classification and did not request, reproduce, or disclose the redacted raw value."), 10, True, color=RED)
    add_response_lines(doc, "Reviewer signature / ID:", 1)
    add_response_lines(doc, "Completion date and time:", 1)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    set_run(p.add_run("RETURN INSTRUCTIONS: "), 9.5, True, color=BLUE)
    set_run(p.add_run("Return the completed handout to the designated review administrator. Do not edit the source packet. The administrator will transcribe and validate the response against the response schema; qualified adjudication and project-owner disposition occur separately."), 9.5)
    doc.add_heading("Administrator use only", level=2)
    option_line(doc, "Transcription status:", ["received", "schema-valid", "returned for clarification", "ready for qualified adjudication"])
    add_response_lines(doc, "Response artifact ID / hash:", 1)


def add_prompt_appendix(doc, designation: str, prompt_entry: dict, prompt_binding: dict, prompt_text: str):
    doc.add_page_break()
    doc.add_heading("Appendix A - Prompt record and markup copy", level=1)
    p = doc.add_paragraph()
    set_run(p.add_run("Purpose. "), 10.5, True, color=NAVY)
    set_run(p.add_run("This appendix identifies and reproduces the exact versioned prompt bound to the specialist candidate execution. Reviewers may mark this copy for clarity, focus, scope, evidence discipline, authority boundaries, and future revision. Markup does not change the active prompt until it is approved, versioned, and rebound through change control."), 10.5)
    metadata = [
        ("Specialist", designation),
        ("Prompt ID", prompt_binding["id"]),
        ("Prompt version", prompt_binding["version"]),
        ("Prompt path", prompt_entry["path"]),
        ("Execution-bound SHA-256", prompt_binding["hash"]),
        ("Binding status", "used in this recorded calibration; hash verified"),
    ]
    table = doc.add_table(rows=len(metadata), cols=2)
    set_table_geometry(table, [2300, 7060])
    for row, (label, value) in zip(table.rows, metadata):
        shade(row.cells[0], LIGHT_BLUE)
        set_run(row.cells[0].paragraphs[0].add_run(label), 9.25, True, color=NAVY)
        set_run(row.cells[1].paragraphs[0].add_run(str(value)), 9.25)
    doc.add_heading("Full prompt - line-numbered markup copy", level=2)
    p = doc.add_paragraph()
    set_run(p.add_run("The text below is a complete UTF-8 reproduction of the bound prompt. Line numbers are presentation aids and are not part of the prompt."), 9.5, italic=True, color=GRAY)
    prompt_table = doc.add_table(rows=1, cols=2)
    set_table_geometry(prompt_table, [700, 8660])
    prompt_table.style = "Table Grid"
    set_repeat_table_header(prompt_table.rows[0])
    for cell, label in zip(prompt_table.rows[0].cells, ("Line", "Prompt text")):
        shade(cell, LIGHT_BLUE)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        set_run(paragraph.add_run(label), 8.5, True, color=NAVY, font="Consolas")
    for line_number, line in enumerate(prompt_text.splitlines(), 1):
        cells = prompt_table.add_row().cells
        cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for cell in cells:
            cell_margins(cell, top=25, start=60, bottom=25, end=60)
            cell.paragraphs[0].paragraph_format.space_before = Pt(0)
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            cell.paragraphs[0].paragraph_format.line_spacing = 1.0
        set_run(cells[0].paragraphs[0].add_run(str(line_number)), 8, color=GRAY, font="Consolas")
        set_run(cells[1].paragraphs[0].add_run(line), 8, color=BLACK, font="Consolas")


def build_document(packet_path: Path, profile: dict, packet: dict, prompt_entry: dict, prompt_binding: dict, prompt_text: str, output: Path):
    doc = Document()
    configure_styles(doc)
    configure_page(doc, packet["candidate"]["designation"], packet["classification"])
    props = doc.core_properties
    props.author = "thomasverburgt"
    props.last_modified_by = "thomasverburgt"
    props.title = f"{packet['candidate']['designation']} Specialist Human Shadow Review Handout"
    props.subject = "ADR-0039 comparison-only specialist candidate review"
    props.keywords = "ADR-0039, specialist, human review, calibration"
    add_title_block(doc, profile, packet)
    add_instructions_and_focus(doc, profile, packet)
    add_evidence_map(doc, packet)
    for index, item in enumerate(packet["items"], 1):
        add_item_review(doc, item, index)
    add_overall_review(doc, profile, packet)
    add_prompt_appendix(doc, packet["candidate"]["designation"], prompt_entry, prompt_binding, prompt_text)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> int:
    catalog = load(ROOT / "agents/focus-profiles/catalog.json")
    prompt_manifest = load(ROOT / "appendices/prompt-templates/candidates/manifest-design-0.2.0.json")
    profiles = {entry["designation"]: load(ROOT / entry["path"]) for entry in catalog["profiles"] if entry["designation"].startswith("SPEC-")}
    packets = {}
    for root in PACKET_ROOTS:
        for path in root.glob("spec-*/human-shadow-review-packet.json"):
            packet = load(path)
            packets[packet["candidate"]["designation"]] = (path, packet)
    if set(packets) != set(profiles):
        raise SystemExit(f"packet/profile mismatch: packets={len(packets)} profiles={len(profiles)} missing={sorted(set(profiles)-set(packets))}")
    OUTPUT.mkdir(parents=True, exist_ok=True)
    manifest = {"schema_version":"1.0.0","governing_decision":"ADR-0039","generated_at":"2026-08-04","authority_effect":"review_handout_only","handouts":[]}
    for designation in sorted(packets):
        path, packet = packets[designation]
        candidate = load(path.parent / "model-candidate.artifact.json")
        prompt_binding = candidate["execution"]["execution_bindings"]["prompt"]
        prompt_entry = prompt_manifest["prompts"][designation]
        prompt_path = ROOT / prompt_entry["path"]
        prompt_bytes = prompt_path.read_bytes()
        prompt_digest = hashlib.sha256(prompt_bytes).hexdigest()
        if prompt_binding["version"] != prompt_manifest["prompt_version"]:
            raise SystemExit(f"prompt version mismatch for {designation}")
        if prompt_digest != prompt_entry["sha256"] or prompt_binding["hash"] != "sha256:" + prompt_digest:
            raise SystemExit(f"prompt hash mismatch for {designation}")
        prompt_text = prompt_bytes.decode("utf-8")
        slug = designation.lower().replace("_", "-")
        output = OUTPUT / f"{slug}-human-shadow-review-handout.docx"
        build_document(path, profiles[designation], packet, prompt_entry, prompt_binding, prompt_text, output)
        digest = hashlib.sha256(output.read_bytes()).hexdigest()
        manifest["handouts"].append({"designation":designation,"display_name":profiles[designation]["agent"]["display_name"],"classification":packet["classification"],"packet_id":packet["packet_id"],"packet_hash":packet["packet_hash"],"source_packet":str(path.relative_to(ROOT)).replace("\\","/"),"prompt":{"id":prompt_binding["id"],"version":prompt_binding["version"],"path":prompt_entry["path"],"sha256":prompt_binding["hash"]},"document":str(output.relative_to(ROOT)).replace("\\","/"),"sha256":"sha256:"+digest})
    manifest["handout_count"] = len(manifest["handouts"])
    (OUTPUT / "manifest.json").write_text(json.dumps(manifest, indent=2, sort_keys=True)+"\n", encoding="utf-8")
    print(json.dumps({"output":str(OUTPUT),"handouts":len(manifest["handouts"])},sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
